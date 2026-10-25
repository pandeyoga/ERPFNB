"""Regression: Report Center › Daily Sales — on-screen preview (RPT-DS-05).

The preview endpoint and the Excel export share `daily_sales_rows`, so the preview totals
must equal the TOTAL row of the workbook for the same filters.
"""
import io

from openpyxl import load_workbook

from .conftest import api_get

RANGE = "date_from=2026-01-01&date_to=2026-12-31"


def test_preview_matches_excel_totals(http_client, admin_token):
    r = api_get(http_client, f"/reports/sales/daily-sales?{RANGE}&limit=5", admin_token)
    assert r.status_code == 200, r.text
    body = r.json()
    rows, meta = body["data"], body["meta"]
    assert len(rows) <= 5 and meta["total"] >= len(rows)
    assert all(row["status"] == "validated" for row in rows)

    x = api_get(http_client, f"/reports/sales/daily-sales.xlsx?{RANGE}", admin_token)
    assert x.status_code == 200
    ws = load_workbook(io.BytesIO(x.content)).active
    total_row = next(row for row in ws.iter_rows(values_only=True) if "TOTAL" in row)
    assert round(float(total_row[3]), 2) == meta["grand_total"]
    assert int(total_row[4]) == meta["transaction_count"]


def test_preview_rejects_bad_date(http_client, admin_token):
    r = api_get(http_client, "/reports/sales/daily-sales?date_from=25-10-2026", admin_token)
    assert r.status_code == 422
